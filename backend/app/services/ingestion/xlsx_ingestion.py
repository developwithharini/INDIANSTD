import os
import re
import uuid
import openpyxl
from typing import Dict, Any, List, Tuple
from sqlalchemy.orm import Session
from app.core.config import settings
from app.models.domain import (
    Standard, StandardVersion, StandardAmendment, StandardRelationship, QCO, IngestionRun, SourceRecord
)

# Header Column Normalization Aliases
HEADER_ALIASES = {
    "standard_number": [
        "standard number", "is number", "is no.", "is no", "is code", "indian standard", "standard_no", "std_no"
    ],
    "title": [
        "title", "standard title", "name", "is title", "subject", "title of standard"
    ],
    "scope": [
        "scope", "scope summary", "description", "standard scope", "abstract", "summary"
    ],
    "category": [
        "category", "product category", "standard category", "type of standard", "group", "product group"
    ],
    "domain": [
        "domain", "sector", "industry", "technical area", "department", "technical committee", "committee"
    ],
    "publication_year": [
        "publication year", "year", "published", "publication", "date of publication", "year of publication", "date of publish"
    ],
    "status": [
        "status", "standard status", "current status", "validity"
    ],
    "qco": [
        "qco", "qco mandate", "mandatory qco", "quality control order", "crs mandate", "compulsory registration"
    ]
}

def normalize_header(header: str) -> str:
    """Normalize raw Excel header string to canonical field name if matched."""
    if not header:
        return ""
    clean = str(header).strip().lower()
    
    # Exact match first
    for field, aliases in HEADER_ALIASES.items():
        if clean in aliases:
            return field
            
    # Partial match for specific fields (avoiding generic "standard")
    if "is number" in clean or "standard number" in clean or clean == "standard":
        return "standard_number"
    if "title" in clean:
        return "title"
    if "scope" in clean or "abstract" in clean:
        return "scope"
    if "type of standard" in clean or "category" in clean:
        return "category"
    if "date of publish" in clean or "publication year" in clean:
        return "publication_year"

    return clean.replace(" ", "_")

def make_standard_id(standard_number: str) -> str:
    """Create deterministic clean alphanumeric ID from standard number."""
    clean = re.sub(r'[^A-Za-z0-9]', '_', str(standard_number)).strip('_')
    clean = re.sub(r'_+', '_', clean)
    return clean.upper() or f"STD_{uuid.uuid4().hex[:8]}"

def process_xlsx_ingestion(file_path: str, db: Session, dataset_version: str = None) -> Dict[str, Any]:
    """
    Multi-sheet Excel (.xlsx) ingestion engine using openpyxl.
    Inspects all sheets, normalizes column headers, upserts records to SQLite,
    and returns a structured IngestionReport dict.
    """
    if not dataset_version:
        dataset_version = settings.DATASET_VERSION

    run_id = f"run_{uuid.uuid4().hex[:10]}"
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet_names = wb.sheetnames

    sheets_discovered = len(sheet_names)
    sheets_used = 0
    records_count = 0
    valid_count = 0
    inserted_count = 0
    updated_count = 0
    skipped_count = 0
    failed_count = 0

    file_basename = os.path.basename(file_path)

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            continue

        # Dynamic header row detection (scans top 5 rows)
        header_row_idx = 0
        raw_headers = []
        mapped_headers = []

        for row_candidate_idx in range(min(5, len(rows))):
            candidate_raw = [str(cell).strip() if cell is not None else "" for cell in rows[row_candidate_idx]]
            candidate_mapped = [normalize_header(h) for h in candidate_raw]
            if "standard_number" in candidate_mapped or "title" in candidate_mapped:
                header_row_idx = row_candidate_idx
                raw_headers = candidate_raw
                mapped_headers = candidate_mapped
                break
        
        if not mapped_headers:
            raw_headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
            mapped_headers = [normalize_header(h) for h in raw_headers]

        sheets_used += 1

        for row_idx, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
            if not row or not any(row):
                continue

            records_count += 1
            row_dict = {}
            extra_metadata = {}

            for col_idx, cell_val in enumerate(row):
                if col_idx < len(mapped_headers):
                    key = mapped_headers[col_idx]
                    val = str(cell_val).strip() if cell_val is not None else ""
                    if key in HEADER_ALIASES:
                        row_dict[key] = val
                    elif raw_headers[col_idx]:
                        extra_metadata[raw_headers[col_idx]] = val

            std_num = row_dict.get("standard_number", "").strip()
            title = row_dict.get("title", "").strip()

            if not std_num and not title:
                skipped_count += 1
                continue

            if not std_num and title:
                # Synthesize standard number if missing
                std_num = f"IS {title[:20].upper()}"

            valid_count += 1
            std_id = make_standard_id(std_num)

            # Parse publication year
            pub_year = None
            year_str = row_dict.get("publication_year", "")
            year_match = re.search(r'\b(19\d\d|20\d\d)\b', year_str or std_num)
            if year_match:
                pub_year = int(year_match.group(1))

            status = row_dict.get("status", "CURRENT").upper()
            if "SUPERSED" in status or "WITHDRAW" in status:
                status = "SUPERSEDED"
            else:
                status = "CURRENT"

            # Check existing standard in DB
            existing = db.query(Standard).filter(Standard.id == std_id).first()

            if existing:
                existing.standard_number = std_num
                existing.title = title or existing.title
                if row_dict.get("scope"):
                    existing.scope = row_dict.get("scope")
                if row_dict.get("category"):
                    existing.category = row_dict.get("category")
                if row_dict.get("domain"):
                    existing.domain = row_dict.get("domain")
                if pub_year:
                    existing.publication_year = pub_year
                existing.status = status
                existing.extra_metadata = extra_metadata
                updated_count += 1
            else:
                new_std = Standard(
                    id=std_id,
                    standard_number=std_num,
                    title=title or "Indian Standard Specification",
                    scope=row_dict.get("scope", "Scope details preserved from official BIS standards corpus."),
                    category=row_dict.get("category", "General Engineering"),
                    domain=row_dict.get("domain", "Production & General Engineering"),
                    ministry="Bureau of Indian Standards",
                    publication_year=pub_year or 2020,
                    status=status,
                    source_url="",
                    source_file=file_basename,
                    source_sheet=sheet_name,
                    source_row=row_idx,
                    source_type="BIS_PUBLISHED_SNAPSHOT",
                    verified=True,
                    extra_metadata=extra_metadata
                )
                db.add(new_std)
                inserted_count += 1

            # Log source record for provenance tracking
            source_rec = SourceRecord(
                ingestion_run_id=run_id,
                source_file=file_basename,
                source_sheet=sheet_name,
                source_row=row_idx,
                standard_number=std_num,
                raw_payload_json=row_dict,
                status="PROCESSED"
            )
            db.add(source_rec)

    db.commit()

    # Log IngestionRun record
    run_record = IngestionRun(
        id=run_id,
        source_file=file_basename,
        sheets_discovered=sheets_discovered,
        sheets_used=sheets_used,
        records_count=records_count,
        valid_count=valid_count,
        inserted_count=inserted_count,
        updated_count=updated_count,
        skipped_count=skipped_count,
        failed_count=failed_count,
        dataset_version=dataset_version
    )
    db.add(run_record)
    db.commit()

    return {
        "ingestion_run_id": run_id,
        "source_file": file_basename,
        "sheets_discovered": sheets_discovered,
        "sheets_used": sheets_used,
        "records_count": records_count,
        "valid_count": valid_count,
        "inserted_count": inserted_count,
        "updated_count": updated_count,
        "skipped_count": skipped_count,
        "failed_count": failed_count,
        "dataset_version": dataset_version,
        "created_at": run_record.created_at.isoformat()
    }
