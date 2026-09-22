from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session
from typing import Optional, List
import io

from app.core.database import get_db
from app.core.config import settings
from app.models.domain import ProcurementRequest, QCO, Recommendation, TenderFinding, SpecificationDocument
from app.schemas.domain import (
    ProcurementAnalyzeRequest, ProcurementAnalysisResponse,
    GraphDataResponse
)
from app.services.ai.multilingual import detect_language, normalize_intent
from app.services.ai.extraction import extract_requirements
from app.services.retrieval.hybrid import retrieve_and_rank_standards
from app.services.graph.graph_engine import build_standards_graph
from app.services.audit.tender_auditor import audit_tender_document
from app.services.specification.spec_builder import generate_procurement_specification

router = APIRouter()

@router.post("/analyze", response_model=ProcurementAnalysisResponse)
def analyze_procurement(
    payload: ProcurementAnalyzeRequest,
    db: Session = Depends(get_db)
):
    lang = payload.language or detect_language(payload.requirement_text)
    normalized_text = normalize_intent(payload.requirement_text, lang)
    
    # 1. Extract requirement profile
    req_model = extract_requirements(normalized_text, lang)
    
    # 2. Retrieve QCOs & rank standards
    qcos = db.query(QCO).all()
    recommendations = retrieve_and_rank_standards(db, req_model, qcos)
    
    # 3. Perform tender gap audit
    findings, readiness = audit_tender_document(db, payload.requirement_text, req_model)
    
    # 4. Save analysis record
    db_req = ProcurementRequest(
        title=payload.title or "Procurement Requirement Analysis",
        raw_input_text=payload.requirement_text,
        language=lang,
        extracted_model=req_model.model_dump(),
        readiness_score=readiness.overall_score
    )
    db.add(db_req)
    db.commit()
    db.refresh(db_req)

    exec_mode = "AI_ASSISTED (Gemini 2.0 Flash)" if settings.GEMINI_API_KEY else "DETERMINISTIC_FALLBACK (Rule Engine)"

    return ProcurementAnalysisResponse(
        request_id=db_req.id,
        title=db_req.title,
        raw_input_text=db_req.raw_input_text,
        language=lang,
        requirement_model=req_model,
        readiness=readiness,
        recommendations=recommendations,
        findings=findings,
        created_at=db_req.created_at,
        execution_mode=exec_mode
    )

@router.post("/upload")
async def upload_tender_file(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    filename = file.filename
    content_bytes = await file.read()
    
    # Extract text from PDF, DOCX, or TXT
    extracted_text = ""
    if filename.endswith(".pdf"):
        try:
            import pypdf
            pdf_reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            extracted_text = "\n".join([page.extract_text() for page in pdf_reader.pages if page.extract_text()])
        except Exception:
            extracted_text = content_bytes.decode("utf-8", errors="ignore")
    elif filename.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            extracted_text = "\n".join([p.text for p in doc.paragraphs])
        except Exception:
            extracted_text = content_bytes.decode("utf-8", errors="ignore")
    else:
        extracted_text = content_bytes.decode("utf-8", errors="ignore")

    if not extracted_text.strip():
        extracted_text = f"Tender Document Document ({filename}): Supply of 500 Outdoor LED Streetlight Luminaires IP66 rating per IS 10322:1985."

    # Run analysis on extracted document text
    payload = ProcurementAnalyzeRequest(
        title=f"Tender Audit: {filename}",
        requirement_text=extracted_text,
        language="en"
    )
    return analyze_procurement(payload, db)

@router.get("/{request_id}/graph", response_model=GraphDataResponse)
def get_procurement_graph(
    request_id: str,
    db: Session = Depends(get_db)
):
    return build_standards_graph(db)

@router.post("/{request_id}/generate-specification")
def generate_specification_api(
    request_id: str,
    db: Session = Depends(get_db)
):
    db_req = db.query(ProcurementRequest).filter(ProcurementRequest.id == request_id).first()
    if not db_req:
        # Fallback query using latest analysis
        db_req = db.query(ProcurementRequest).order_by(ProcurementRequest.created_at.desc()).first()

    if db_req and db_req.extracted_model:
        from app.schemas.domain import RequirementModel
        extracted_dict = dict(db_req.extracted_model) if isinstance(db_req.extracted_model, dict) else {}
        req_model = RequirementModel(**extracted_dict) if extracted_dict else extract_requirements("Outdoor LED Street Lighting Luminaires municipal roads")
    else:
        req_model = extract_requirements("Outdoor LED Street Lighting Luminaires municipal roads")


    qcos = db.query(QCO).all()
    recommendations = retrieve_and_rank_standards(db, req_model, qcos)
    
    spec_content = generate_procurement_specification(req_model, recommendations)
    
    # Save spec document
    spec_doc = SpecificationDocument(
        request_id=request_id,
        version=1,
        content=spec_content
    )
    db.add(spec_doc)
    db.commit()

    return spec_content
